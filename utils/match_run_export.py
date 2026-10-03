"""
utils/match_run_export.py

Builds a formatted .xlsx for ONE past match run, using only what's stored
in the match_runs table (matches/waitlist/dropped_mcps as IDs + metadata) —
not the original YP/MCP person records, which aren't persisted anywhere.

This is deliberately separate from utils/excel_export.build_results_workbook,
which expects the full in-memory YP/MCP objects (names, addresses, etc.) and
is only ever called right after a fresh /match/run while that data is still
in `state`. A past run has no guarantee the person who uploaded the original
YP/MCP sheets is still the one re-downloading it, or that `state` still
holds the same data — so this reads straight from the stored run and never
touches `state`. If you want historical exports with full person details
(names/addresses) rather than just IDs, that requires snapshotting the YP/
MCP records into match_runs at save time — a bigger schema change; this
module does not attempt that.
"""

import logging
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)


def _write_sheet(ws, headers: list[str], rows: list[list[Any]]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for row in rows:
        ws.append(row)

    for col_idx, header in enumerate(headers, start=1):
        longest = max(
            [len(str(header))] + [len(str(row[col_idx - 1])) for row in rows],
            default=len(str(header)),
        )
        ws.column_dimensions[get_column_letter(col_idx)].width = min(longest + 2, 60)

    ws.freeze_panes = "A2"


def build_match_run_workbook(run: dict[str, Any], output_path: str) -> None:
    """
    run: the dict returned by match_run_store.get_match_run() — must include
    id, created_at, actor_email, hop_limit, match_cap, shortlist_size,
    matched_count, waitlisted_count, dropped_mcp_count, matches, waitlist,
    dropped_mcps.

    Raises RuntimeError (not the raw underlying exception) if the run dict
    is missing an expected key/shape, or if writing the workbook fails —
    e.g. output_path's directory doesn't exist or isn't writable.
    """
    try:
        wb = Workbook()

        summary_ws = wb.active
        summary_ws.title = "Summary"
        summary_rows = [
            ["Run ID", run["id"]],
            ["Run at", str(run["created_at"])],
            ["Run by", run.get("actor_email") or run.get("actor_name") or run.get("actor_id")],
            ["Hop limit", run["hop_limit"]],
            ["Match cap", run["match_cap"] if run["match_cap"] is not None else "(none)"],
            ["Shortlist size", run["shortlist_size"]],
            ["Matched count", run["matched_count"]],
            ["Waitlisted count", run["waitlisted_count"]],
            ["Dropped MCP count", run["dropped_mcp_count"]],
        ]
        for row in summary_rows:
            summary_ws.append(row)
        for cell in summary_ws["A"]:
            cell.font = Font(bold=True)
        summary_ws.column_dimensions["A"].width = 20
        summary_ws.column_dimensions["B"].width = 40

        matches_ws = wb.create_sheet("Matches")
        match_rows = [
            [m["yp_id"], m["mcp_id"], m.get("landmark"), m.get("travel_time"), m.get("round")]
            for m in run["matches"]
        ]
        _write_sheet(matches_ws, ["YP ID", "MCP ID", "Landmark", "Travel Time", "Round"], match_rows)

        waitlist_ws = wb.create_sheet("Waitlist")
        waitlist_rows = [[w["yp_id"], w.get("reason")] for w in run["waitlist"]]
        _write_sheet(waitlist_ws, ["YP ID", "Reason"], waitlist_rows)

        dropped_ws = wb.create_sheet("Dropped MCPs")
        dropped_rows = [[d["mcp_id"], d.get("reason")] for d in run["dropped_mcps"]]
        _write_sheet(dropped_ws, ["MCP ID", "Reason"], dropped_rows)

        wb.save(output_path)
    except (KeyError, TypeError, OSError) as e:
        logger.exception(
            "match_run_export: failed to build workbook for run id=%s",
            run.get("id") if isinstance(run, dict) else "?",
        )
        raise RuntimeError(f"Could not build export workbook: {e}") from e